import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from engine.utils.pdf import PDFConversionError, convert_docx_to_pdf


class PDFConversionTest(unittest.TestCase):
    def test_missing_libreoffice_reports_clear_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "sample.docx"
            docx_path.write_bytes(b"docx")

            with patch("engine.utils.pdf.which", return_value=None):
                with self.assertRaisesRegex(PDFConversionError, "LibreOffice executable not found"):
                    convert_docx_to_pdf(docx_path, Path(tmp) / "pdf")

    def test_successful_conversion_returns_expected_pdf_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            docx_path = root / "sample.docx"
            pdf_dir = root / "pdf"
            docx_path.write_bytes(b"docx")

            def fake_run(cmd, check, text, stdout, stderr):
                pdf_dir.mkdir(parents=True, exist_ok=True)
                (pdf_dir / "sample.pdf").write_bytes(b"%PDF-1.7")
                return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")

            with patch("engine.utils.pdf.which", return_value="/usr/bin/soffice"):
                with patch("engine.utils.pdf.subprocess.run", side_effect=fake_run):
                    pdf_path = convert_docx_to_pdf(docx_path, pdf_dir)

            self.assertEqual(pdf_path, pdf_dir / "sample.pdf")
            self.assertTrue(pdf_path.exists())

    def test_failed_libreoffice_uses_fallback_pdf_conversion(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            docx_path = root / "sample.docx"
            pdf_dir = root / "pdf"
            document = Document()
            document.add_heading("Document de test", level=1)
            document.add_paragraph("Contenu validé.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Référence"
            table.cell(0, 1).text = "Document"
            table.cell(1, 0).text = "CEVA-RHYAD-200-F01-DT01"
            table.cell(1, 1).text = "Données Techniques"
            document.save(docx_path)

            def fake_run(cmd, check, text, stdout, stderr):
                return subprocess.CompletedProcess(cmd, 134, stdout="", stderr="Abort trap: 6")

            with patch("engine.utils.pdf.which", return_value="/usr/bin/soffice"):
                with patch("engine.utils.pdf.subprocess.run", side_effect=fake_run):
                    pdf_path = convert_docx_to_pdf(docx_path, pdf_dir)

            self.assertEqual(pdf_path, pdf_dir / "sample.pdf")
            self.assertTrue(pdf_path.exists())
            self.assertTrue(pdf_path.read_bytes().startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
