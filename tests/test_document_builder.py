import tempfile
import unittest
from pathlib import Path

from docx import Document

from engine.core.document_builder import build_document, validate_document_config


class DocumentBuilderTest(unittest.TestCase):
    def test_build_document_with_project_metadata_and_optional_missing_logo(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            yaml_path = root / "F99.yaml"
            yaml_path.write_text(
                "\n".join(
                    [
                        "reference: TEST-F99",
                        "title: Test document",
                        "subtitle: Revision T",
                        "chapters:",
                        "  - title: 1. Purpose",
                        "    text: Test body.",
                        "  - title: 2. References",
                        "    bullets:",
                        "      - REF-001",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = root / "output" / "docx" / "TEST-F99.docx"
            project_config = {
                "project": {
                    "name": "Test Project",
                    "client": "Test Client",
                    "location": "Test Location",
                },
                "document": {
                    "revision": "T1",
                    "status": "DRAFT",
                    "confidentiality": "INTERNAL",
                },
                "authors": {"author": "Test Author", "company": "Test Company"},
                "branding": {"logo_ceva": "assets/logos/missing.png"},
            }

            build_document(
                yaml_path,
                output_path,
                project_config=project_config,
                project_root=root,
            )

            self.assertTrue(output_path.exists())
            generated = Document(output_path)
            paragraph_text = "\n".join(p.text for p in generated.paragraphs)
            table_text = "\n".join(
                cell.text for table in generated.tables for row in table.rows for cell in row.cells
            )

            self.assertIn("TEST-F99", paragraph_text)
            self.assertIn("Test document", paragraph_text)
            self.assertIn("REF-001", paragraph_text)
            self.assertIn("Test Project", table_text)
            self.assertIn("DRAFT", table_text)

    def test_validate_document_config_requires_chapters(self):
        with self.assertRaisesRegex(ValueError, "chapters"):
            validate_document_config(
                {"reference": "TEST-F99", "title": "Test document"},
                Path("F99.yaml"),
            )


if __name__ == "__main__":
    unittest.main()
