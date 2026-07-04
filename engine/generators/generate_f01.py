from docx import Document
from engine.styles.ceva_style import setup_document, add_title

def generate(output_path):
    doc = Document()
    setup_document(doc)

    add_title(
        doc,
        "RHYAD-F01",
        "Gestion des actions",
        "Procédure projet - Révision A1.0"
    )

    chapters = [
        ("1. Objet", "Définir la méthode unique de création, attribution, suivi, mise à jour et clôture des actions du projet RHYAD."),
        ("2. Domaine d'application", "Applicable à la MOA, AMO, MOE, architectes, ingénieries, entreprises, fournisseurs et contrôleurs."),
        ("3. Documents de référence", "DT-000, REG-ACT, REG-DEC, F02, F03, F04, F05, ISO 9001, ISO 21502."),
        ("4. Définitions", "Action, responsable, échéance, clôture, escalade, preuve documentaire."),
        ("5. Principes généraux", "Toute action doit avoir un identifiant unique, un responsable unique, une échéance, une priorité, un statut et une preuve de clôture."),
        ("6. Typologie des actions", "TECH, DOC, DEC, RFI, RISK, QUALITY, PROCUREMENT, SITE, HSE, CLIENT."),
        ("7. Cycle de vie", "Identification → Création → Validation → Affectation → En cours → Suivi → Terminée → Vérification → Clôture."),
        ("8. Création d'une action", "Chaque action est saisie dans REG-ACT avec les champs obligatoires définis par la gouvernance documentaire RHYAD."),
        ("9. Priorisation", "P1 Critique : 24 h ; P2 Haute : 5 jours ; P3 Normale : 15 jours ; P4 Faible : selon planning."),
        ("10. Responsabilités", "Un responsable unique est désigné. Une matrice RACI est utilisée lorsque plusieurs disciplines sont impliquées."),
        ("11. Suivi des actions", "Le registre des actions est mis à jour au minimum une fois par semaine et revu en réunion de coordination."),
        ("12. Clôture", "Une action ne peut être clôturée qu'après validation du livrable attendu et enregistrement de la preuve documentaire."),
        ("13. KPI", "Actions ouvertes, actions clôturées, taux de clôture à l'échéance, actions critiques en retard, âge moyen des actions."),
        ("14. Interfaces", "Le processus est lié aux décisions, RFI, interfaces, risques, modifications et non-conformités."),
        ("15. Annexes", "Modèle de registre, RACI, tableau de bord KPI, exemple de fiche action.")
    ]

    for heading, body in chapters:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)

    doc.save(output_path)
