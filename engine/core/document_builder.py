from pathlib import Path
from datetime import date
import yaml
from docx import Document

from engine.styles.ceva_style import (
    add_cover_page,
    add_data_table,
    add_document_control_table,
    add_table_of_contents,
    setup_document_shell,
)


def _require_mapping(data, label):
    if not isinstance(data, dict):
        raise ValueError(f"{label} must be a YAML mapping.")


def _require_string(data, key, label):
    value = data.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label}.{key} must be a non-empty string.")
    return value


def validate_document_config(data, yaml_path: Path):
    label = str(yaml_path)
    _require_mapping(data, label)

    _require_string(data, "reference", label)
    _require_string(data, "title", label)

    subtitle = data.get("subtitle")
    if subtitle is not None and not isinstance(subtitle, str):
        raise ValueError(f"{label}.subtitle must be a string when provided.")

    chapters = data.get("chapters")
    if not isinstance(chapters, list) or not chapters:
        raise ValueError(f"{label}.chapters must be a non-empty list.")

    for index, chapter in enumerate(chapters, start=1):
        chapter_label = f"{label}.chapters[{index}]"
        _require_mapping(chapter, chapter_label)
        _require_string(chapter, "title", chapter_label)

        text = chapter.get("text")
        if text is not None and not isinstance(text, str):
            raise ValueError(f"{chapter_label}.text must be a string when provided.")

        bullets = chapter.get("bullets")
        if bullets is not None:
            if not isinstance(bullets, list):
                raise ValueError(f"{chapter_label}.bullets must be a list when provided.")
            for item_index, item in enumerate(bullets, start=1):
                if not isinstance(item, str) or not item.strip():
                    raise ValueError(
                        f"{chapter_label}.bullets[{item_index}] must be a non-empty string."
                    )

        tables = chapter.get("tables")
        if tables is not None:
            if not isinstance(tables, list) or not tables:
                raise ValueError(f"{chapter_label}.tables must be a non-empty list when provided.")
            for table_index, table in enumerate(tables, start=1):
                table_label = f"{chapter_label}.tables[{table_index}]"
                _require_mapping(table, table_label)

                title = table.get("title")
                if title is not None and not isinstance(title, str):
                    raise ValueError(f"{table_label}.title must be a string when provided.")

                headers = table.get("headers")
                if not isinstance(headers, list) or not headers:
                    raise ValueError(f"{table_label}.headers must be a non-empty list.")
                for header_index, header in enumerate(headers, start=1):
                    if not isinstance(header, str) or not header.strip():
                        raise ValueError(
                            f"{table_label}.headers[{header_index}] must be a non-empty string."
                        )

                rows = table.get("rows")
                if not isinstance(rows, list) or not rows:
                    raise ValueError(f"{table_label}.rows must be a non-empty list.")
                for row_index, row in enumerate(rows, start=1):
                    row_label = f"{table_label}.rows[{row_index}]"
                    if not isinstance(row, list):
                        raise ValueError(f"{row_label} must be a list.")
                    if len(row) != len(headers):
                        raise ValueError(f"{row_label} must contain {len(headers)} values.")


def _resolve_optional_path(path_value, project_root):
    if not path_value:
        return None

    path = Path(path_value)
    if not path.is_absolute() and project_root is not None:
        path = Path(project_root) / path

    return path if path.exists() else None


def _metadata(data, project_config):
    project_config = project_config or {}
    project = project_config.get("project") or {}
    document = project_config.get("document") or {}
    authors = project_config.get("authors") or {}
    approval = project_config.get("approval") or {}

    return {
        "reference": data["reference"],
        "title": data["title"],
        "subtitle": data.get("subtitle", ""),
        "project_name": project.get("name", ""),
        "client": project.get("client", ""),
        "location": project.get("location", ""),
        "revision": data.get("revision", document.get("revision", "")),
        "status": data.get("status", document.get("status", "")),
        "confidentiality": document.get("confidentiality", ""),
        "author": authors.get("author", ""),
        "checker": approval.get("checker", ""),
        "approver": approval.get("approver", ""),
        "generated_date": date.today().isoformat(),
    }


def build_document(yaml_path: Path, output_path: Path, project_config=None, project_root=None):
    with open(yaml_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)

    validate_document_config(data, yaml_path)

    doc = Document()

    branding = (project_config or {}).get("branding", {})
    logo_path = _resolve_optional_path(branding.get("logo_ceva"), project_root)
    metadata = _metadata(data, project_config or {})

    setup_document_shell(doc, metadata, logo_path=logo_path)
    add_cover_page(doc, metadata, logo_path=logo_path)
    add_document_control_table(doc, metadata)
    add_table_of_contents(doc)

    for chapter in data["chapters"]:
        doc.add_heading(chapter["title"], level=1)

        if "text" in chapter:
            doc.add_paragraph(chapter["text"])

        if "bullets" in chapter:
            for item in chapter["bullets"]:
                doc.add_paragraph(item, style="List Bullet")

        if "tables" in chapter:
            for table in chapter["tables"]:
                add_data_table(doc, table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
